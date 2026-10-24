"""Document ingestion: parse PDF/DOCX/PPTX/TXT and chunk text."""
from __future__ import annotations

from pathlib import Path


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for i, page in enumerate(reader.pages):
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        if txt.strip():
            parts.append(f"[Page {i + 1}]\n{txt}")
    return "\n\n".join(parts)


def _read_docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _read_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides):
        lines = [f"[Slide {i + 1}]"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs)
                    if text.strip():
                        lines.append(text)
        if len(lines) > 1:
            parts.append("\n".join(lines))
    return "\n\n".join(parts)


def _read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


READERS = {
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".doc": _read_docx,
    ".pptx": _read_pptx,
    ".ppt": _read_pptx,
    ".txt": _read_txt,
    ".md": _read_txt,
}


def extract_text(path: str | Path) -> str:
    path = Path(path)
    reader = READERS.get(path.suffix.lower())
    if reader is None:
        raise ValueError(f"Unsupported file type: {path.suffix}")
    return reader(path)


def chunk_text(
    text: str, chunk_size: int = 1200, overlap: int = 200
) -> list[str]:
    """Split text into overlapping word-based chunks."""
    text = " ".join(text.split())
    if not text:
        return []
    words = text.split(" ")
    chunks: list[str] = []
    step = max(1, chunk_size - overlap)
    for start in range(0, len(words), step):
        chunk = " ".join(words[start : start + chunk_size])
        if chunk.strip():
            chunks.append(chunk)
        if start + chunk_size >= len(words):
            break
    return chunks